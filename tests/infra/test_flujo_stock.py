"""Pruebas de integración end-to-end del flujo de Stock sobre archivos .xlsx
sintéticos: Módulo A (clasificación vía columnas auxiliares opcionales), Módulo B
(ubicación por defecto), Módulo C (validación) e idempotencia de la transferencia."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from airandes_migrador.core.stock import io_excel
from airandes_migrador.core.stock.config import ConfigStock
from airandes_migrador.core.stock.modelo import NOMBRES_COLUMNAS_EXCEL
from airandes_migrador.infra.db import AuditoriaDB

COLUMNAS_ENCABEZADO = NOMBRES_COLUMNAS_EXCEL + ["FAMILIA_CLASEX", "CANTIDAD_NEGATIVA"]


def _crear_workbook_trabajo(ruta: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "STOCK"
    ws.append(COLUMNAS_ENCABEZADO)

    # Fila 1: ya clasificada, sin ubicación (Módulo B debe inyectar "POR UBICAR").
    ws.append(
        ["TORNILLO AN3-4A", "AN3-4A", None, None, "Consumibles", "07", "Z", None, "Boeing",
         None, "AIRANDES", "DEPOSITO CENTRAL", None, None, None, 10, "unid.", "USD", 1.5,
         None, None, False]
    )
    # Fila 2: sin clasificar, se resuelve vía FAMILIA_CLASEX (Módulo A).
    ws.append(
        ["TORQUIMETRO PAÑOL", "PN-9", "SN-9", None, None, None, None, None, "Britool",
         "Disponible", "AIRANDES", "DEPOSITO CENTRAL", "ESTANTE B2", None, None, 1, "UND",
         "USD", 200.0, None, "HERRAMIENTAS", False]
    )
    # Fila 3: cantidad negativa -> debe quedar pendiente, nunca autotransferible.
    ws.append(
        ["FILTRO ACEITE", "PN-7", None, None, "Consumibles", "07", "Z", None, "Champion",
         None, "AIRANDES", "DEPOSITO CENTRAL", "ESTANTE C3", None, None, -3, "UND", "USD",
         20.0, None, None, True]
    )

    wb.create_sheet("Trazabilidad_ClaseX")["A1"] = "no tocar"
    wb.save(ruta)


def _crear_workbook_maestro(ruta: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "STOCK"
    ws.append(NOMBRES_COLUMNAS_EXCEL)
    # columnas vacías extra hasta la Z, tal como el maestro real: no deben afectar la lectura.
    for col in range(len(NOMBRES_COLUMNAS_EXCEL) + 1, 27):
        ws.cell(row=1, column=col, value=None)
    wb.save(ruta)


@pytest.fixture
def config_stock_real() -> ConfigStock:
    from airandes_migrador.core.stock.config import cargar_config_stock

    directorio_config = Path(__file__).resolve().parents[2] / "config"
    return cargar_config_stock(directorio_config)


def test_flujo_completo_modulos_a_b_c_e_idempotencia(
    tmp_path: Path, config_stock_real: ConfigStock
):
    ruta_trabajo = tmp_path / "01_Stock_PreCarga_Trabajo.xlsx"
    ruta_maestro = tmp_path / "STOCK_CAVOK.xlsx"
    _crear_workbook_trabajo(ruta_trabajo)
    _crear_workbook_maestro(ruta_maestro)

    evaluadas, resumen, modulos = io_excel.generar_matriz_validacion(
        ruta_trabajo, ruta_trabajo, config_stock_real
    )
    assert resumen.total == 3
    assert modulos.ubicaciones_inyectadas == 1  # fila 1
    assert resumen.aprobadas == 2  # filas 1 y 2
    assert resumen.pendientes == 1  # fila 3 (cantidad negativa)

    fila1 = evaluadas[0]
    assert fila1.fila.ubicacion == "POR UBICAR"
    assert fila1.fila.unidad == "UND"  # homologación "unid." -> "UND"

    fila2 = evaluadas[1]
    assert fila2.fila.categoria == "Herramientas"  # Módulo A vía FAMILIA_CLASEX
    assert fila2.fila.tipo == "08"
    assert fila2.fila.xyz == "Y"

    fila3 = evaluadas[2]
    assert not fila3.resultado.es_transferible

    wb_verif = openpyxl.load_workbook(ruta_trabajo)
    assert wb_verif["Trazabilidad_ClaseX"]["A1"].value == "no tocar"

    audit_db = AuditoriaDB(tmp_path / "auditoria.db")
    resumen_transferencia = io_excel.transferir_aprobadas(
        ruta_trabajo, ruta_maestro, usuario="cargador_test", audit_db=audit_db
    )
    assert resumen_transferencia.transferidas == 2
    assert resumen_transferencia.omitidas_no_aprobadas == 1

    resumen_2 = io_excel.transferir_aprobadas(
        ruta_trabajo, ruta_maestro, usuario="cargador_test", audit_db=audit_db
    )
    assert resumen_2.transferidas == 0
    assert resumen_2.omitidas_ya_transferidas == 2

    wb_maestro = openpyxl.load_workbook(ruta_maestro)
    filas_con_datos = [
        f for f in wb_maestro["STOCK"].iter_rows(min_row=2, values_only=True) if f[0]
    ]
    assert len(filas_con_datos) == 2
