"""Pruebas de integración end-to-end del flujo de Herramientas sobre archivos .xlsx
sintéticos (no los reales del proyecto, para que la suite sea rápida y aislada).
Cubren: no tocar hojas/columnas ajenas, e idempotencia de la transferencia."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pytest

from airandes_migrador.core.herramientas import io_excel
from airandes_migrador.core.herramientas.config import ConfigHerramientas
from airandes_migrador.core.herramientas.modelo import NOMBRES_COLUMNAS_EXCEL
from airandes_migrador.infra.db import AuditoriaDB

FECHA_ACTUAL = date(2026, 9, 3)


def _crear_workbook_trabajo(ruta: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Herramienta"
    ws.append(NOMBRES_COLUMNAS_EXCEL)
    fila_aprobable = [
        "AIRANDES", "San Fernando", "TORQUIMETRO", None, None, "cdi", "PN-1", "SN-001",
        None, None, "UND", "HANGAR N°1", "Disponible", 365, 30, date(2026, 6, 1),
        "JET SERVICE", None, None, None, None, None, None, None, None, None, None,
        None, None, None, None, None, None, None, None, None,
    ]
    fila_incompleta = [
        "AIRANDES", "San Fernando", "PIE DE REY", None, None, "mitutoyo", "PN-2", None,
        None, None, "UND", "HANGAR N°1", "Disponible", 365, 30, date(2026, 1, 1),
        "CEMEC", None, None, None, None, None, None, None, None, None, None,
        None, None, None, None, None, None, None, None, None,
    ]
    ws.append(fila_aprobable)
    ws.append(fila_incompleta)

    hoja_leame = wb.create_sheet("LÉAME")
    hoja_leame["A1"] = "No tocar esta hoja"

    wb.save(ruta)


def _crear_workbook_maestro(ruta: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Herramienta"
    ws.append(NOMBRES_COLUMNAS_EXCEL)
    wb.save(ruta)


@pytest.fixture
def config_herramientas_real() -> ConfigHerramientas:
    from airandes_migrador.core.herramientas.config import cargar_config_herramientas

    directorio_config = Path(__file__).resolve().parents[2] / "config"
    return cargar_config_herramientas(directorio_config)


def test_flujo_completo_no_toca_otras_hojas_y_es_idempotente(
    tmp_path: Path, config_herramientas_real: ConfigHerramientas
):
    ruta_trabajo = tmp_path / "01_Trabajo_Calibradas.xlsx"
    ruta_maestro = tmp_path / "Import_herramienta.xlsx"
    _crear_workbook_trabajo(ruta_trabajo)
    _crear_workbook_maestro(ruta_maestro)

    evaluadas, resumen = io_excel.generar_matriz_validacion(
        ruta_trabajo, ruta_trabajo, config_herramientas_real, FECHA_ACTUAL
    )
    assert resumen.total == 2
    assert resumen.aprobadas == 1
    assert resumen.rechazadas == 1

    # Regla 3: el fabricante en minúsculas se normalizó.
    fila_aprobada = next(e for e in evaluadas if e.resultado.veredicto.value == "APROBADO")
    assert fila_aprobada.fila.fabricante == "CONSOLIDATED DEVICES INC (CDI)"

    # La hoja LÉAME no debe haberse tocado.
    wb_verif = openpyxl.load_workbook(ruta_trabajo)
    assert wb_verif["LÉAME"]["A1"].value == "No tocar esta hoja"

    audit_db = AuditoriaDB(tmp_path / "auditoria.db")
    resumen_transferencia = io_excel.transferir_aprobadas(
        ruta_trabajo, ruta_maestro, usuario="cargador_test", audit_db=audit_db
    )
    assert resumen_transferencia.transferidas == 1
    assert resumen_transferencia.omitidas_no_aprobadas == 1

    wb_maestro = openpyxl.load_workbook(ruta_maestro)
    ws_maestro = wb_maestro["Herramienta"]
    assert ws_maestro.cell(row=2, column=8).value == "SN-001"  # columna SN

    # Segunda corrida: no debe duplicar la fila ya transferida.
    resumen_2 = io_excel.transferir_aprobadas(
        ruta_trabajo, ruta_maestro, usuario="cargador_test", audit_db=audit_db
    )
    assert resumen_2.transferidas == 0
    assert resumen_2.omitidas_ya_transferidas == 1

    wb_maestro_2 = openpyxl.load_workbook(ruta_maestro)
    filas_con_datos = [
        f for f in wb_maestro_2["Herramienta"].iter_rows(min_row=2, values_only=True) if f[7]
    ]
    assert len(filas_con_datos) == 1
