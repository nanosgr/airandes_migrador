"""Colores de semáforo aplicados a filas de las planillas de validación, replicando
el formato condicional que hacían las macros VBA (ROJO / NARANJA / VERDE)."""

from __future__ import annotations

from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from airandes_migrador.core.common.validacion import Semaforo

_COLOR_HEX = {
    Semaforo.ROJO: "FFC7CE",
    Semaforo.NARANJA: "FFE0B2",
    Semaforo.VERDE: "C6EFCE",
}

_RELLENOS = {
    semaforo: PatternFill(start_color=hexa, end_color=hexa, fill_type="solid")
    for semaforo, hexa in _COLOR_HEX.items()
}


def relleno_para(semaforo: Semaforo) -> PatternFill:
    return _RELLENOS[semaforo]


def aplicar_fill_fila(
    hoja: Worksheet, fila_excel: int, columna_desde: int, columna_hasta: int, semaforo: Semaforo
) -> None:
    """Pinta el rango [columna_desde, columna_hasta] (1-indexado, inclusive) de una
    fila con el color correspondiente al semáforo. `fila_excel` es 1-indexado tal
    como lo espera openpyxl (fila 1 = encabezado)."""
    relleno = relleno_para(semaforo)
    for columna in range(columna_desde, columna_hasta + 1):
        hoja.cell(row=fila_excel, column=columna).fill = relleno
