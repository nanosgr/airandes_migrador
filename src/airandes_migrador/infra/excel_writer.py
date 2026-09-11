"""Escritura de hojas Excel celda por celda: solo toca las columnas que la app es
dueña de escribir (auditoría, semáforo, observaciones) y nunca reconstruye la hoja,
para no perder validaciones de datos, formato ni otras hojas del libro."""

from __future__ import annotations

from pathlib import Path

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from airandes_migrador.core.common.excepciones import ErrorEncabezadoExcel


def agregar_columnas_si_faltan(hoja: Worksheet, columnas_nuevas: list[str]) -> dict[str, int]:
    """Asegura que `columnas_nuevas` existan en el encabezado (fila 1), agregándolas
    al final si no están, y devuelve {nombre: índice 1-indexado} de todas ellas."""
    from airandes_migrador.infra.excel_reader import leer_encabezado

    encabezado = leer_encabezado(hoja)
    indice: dict[str, int] = {}
    siguiente_columna = len(encabezado) + 1
    for nombre in columnas_nuevas:
        if nombre in encabezado:
            indice[nombre] = encabezado.index(nombre) + 1
        else:
            hoja.cell(row=1, column=siguiente_columna, value=nombre)
            indice[nombre] = siguiente_columna
            siguiente_columna += 1
    return indice


def escribir_valores_fila(
    hoja: Worksheet, fila_excel: int, indice_columnas: dict[str, int], valores: dict[str, object]
) -> None:
    for nombre, valor in valores.items():
        if nombre not in indice_columnas:
            raise ErrorEncabezadoExcel(f"Columna '{nombre}' no está indexada para escritura.")
        hoja.cell(row=fila_excel, column=indice_columnas[nombre], value=valor)


def primera_fila_libre(hoja: Worksheet, columna_clave: int = 1) -> int:
    """Primera fila (1-indexada) sin valor en `columna_clave`, arrancando después del
    encabezado. Se usa para insertar al final del maestro sin pisar filas existentes."""
    fila = 2
    while hoja.cell(row=fila, column=columna_clave).value not in (None, ""):
        fila += 1
    return fila


def guardar(wb: Workbook, ruta_salida: Path) -> None:
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)
